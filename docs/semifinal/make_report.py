#!/home/xinyuan/anaconda3/envs/numpy1/bin/python
# -*- coding: utf-8 -*-
"""复赛报告 md -> docx（官方模板视觉规范：标题 #1F4D78 / 强调 #0B2545 / 次要 #5A6573 / 表格底纹 #F4F6F9，
中文微软雅黑 + 西文 Calibri；图与图注按锚点插入；附录 A/B/C 合订；页码）
用法: make_report.py <主文.md> <输出.docx>
"""
import re, sys, docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

F = '/home/xinyuan/比赛/赛道三赛题二/复赛/'
FIG1 = '/home/xinyuan/比赛/赛道三赛题二/figs/'
FIG2 = '/data/xinyuan/GOAI_ai4s_env/e44/figs/'
SRC = sys.argv[1] if len(sys.argv) > 1 else '报告_v6_主文.md'
OUT = sys.argv[2] if len(sys.argv) > 2 else '科学发现与环境定义报告_v6.docx'
md = open(F + SRC).read()

# ---------- 官方配色 / 字体 ----------
C_TITLE = RGBColor(0x1F, 0x4D, 0x78)   # 标题深蓝
C_EMPH = RGBColor(0x0B, 0x25, 0x45)    # 强调藏蓝（正文粗体）
C_MINOR = RGBColor(0x5A, 0x65, 0x73)   # 次要文字灰蓝（图注/引注/页码）
C_BODY = RGBColor(0x26, 0x2B, 0x33)    # 正文
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FILL_ZEBRA = 'F4F6F9'
FILL_HEAD = '1F4D78'
BORDER = 'BFC7D1'
EA = '微软雅黑'; EA_HEAD = '微软雅黑'; LATIN = 'Calibri'; MONO = 'Consolas'

d = docx.Document()
for s in d.sections:
    s.top_margin = s.bottom_margin = Cm(2.0); s.left_margin = s.right_margin = Cm(2.2)


def set_fonts(rpr_owner, latin=LATIN, ea=EA):
    """同时设 ascii/hAnsi/cs 与 eastAsia，否则 Word 不认中文字体"""
    rf = rpr_owner.get_or_add_rPr().get_or_add_rFonts()
    for a in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        if rf.get(qn('w:' + a)) is not None: del rf.attrib[qn('w:' + a)]
    rf.set(qn('w:ascii'), latin); rf.set(qn('w:hAnsi'), latin); rf.set(qn('w:cs'), latin)
    rf.set(qn('w:eastAsia'), ea)


st = d.styles['Normal']
st.font.name = LATIN; st.font.size = Pt(10.5); st.font.color.rgb = C_BODY
set_fonts(st.element)
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.28
for name, sz in (('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11.5)):
    hs = d.styles[name]
    hs.font.name = LATIN; hs.font.size = Pt(sz); hs.font.bold = True; hs.font.color.rgb = C_TITLE
    hs.font.italic = False
    set_fonts(hs.element, LATIN, EA_HEAD)
    hs.paragraph_format.keep_with_next = True

# 行内标记：**粗体** 与 `代码`
TOK = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[A-Za-z][^*\n]{0,40}[A-Za-z0-9.]\*)')


def unescape(t):
    return t.replace('\\|', '|')


def add_runs(p, t, size=10.5, bold=False, ea=EA, color=None, italic=False, bold_color=C_EMPH):
    for part in TOK.split(t):
        if not part: continue
        is_b = part.startswith('**') and part.endswith('**')
        is_c = part.startswith('`') and part.endswith('`')
        is_i = (not is_b) and part.startswith('*') and part.endswith('*')
        txt = part[2:-2] if is_b else (part[1:-1] if (is_c or is_i) else part)
        r = p.add_run(unescape(txt))
        r.font.size = Pt(size); r.font.bold = (bold or is_b); r.font.italic = (italic or is_i)
        if is_c:
            set_fonts(r.element, MONO, ea); r.font.size = Pt(max(size - 0.5, 8))
            r.font.color.rgb = C_EMPH if color is None else color
        else:
            set_fonts(r.element, LATIN, ea)
            if is_b and not bold and bold_color is not None: r.font.color.rgb = bold_color
            elif color is not None: r.font.color.rgb = color
    return p


def para(t, size=10.5, bold=False, ea=EA, space=4, align=None, color=None, italic=False, indent=None,
         hanging=None, bold_color=C_EMPH, style=None):
    p = d.add_paragraph(style=style) if style else d.add_paragraph()
    add_runs(p, t, size, bold, ea, color, italic, bold_color)
    p.paragraph_format.space_after = Pt(space); p.paragraph_format.line_spacing = 1.28
    if align is not None: p.alignment = align
    if indent is not None: p.paragraph_format.left_indent = Cm(indent)
    if hanging is not None: p.paragraph_format.first_line_indent = Cm(-hanging)
    return p


def heading(t, lvl, space_before=None):
    style = {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3'}[lvl]
    size = {1: 16, 2: 12.5, 3: 11.5}[lvl]
    p = para(t, size, True, EA_HEAD, {1: 8, 2: 6, 3: 4}[lvl], color=C_TITLE, bold_color=None, style=style)
    p.paragraph_format.space_before = Pt(space_before if space_before is not None else {1: 14, 2: 10, 3: 6}[lvl])
    p.paragraph_format.keep_with_next = True
    return p


def _pbdr(p, color=FILL_HEAD, sz='18', fill=None):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr'); left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), sz); left.set(qn('w:space'), '8'); left.set(qn('w:color'), color)
    bdr.append(left)
    SUCC = ('w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct',
            'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
            'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
            'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
    ppr.insert_element_before(bdr, *SUCC)
    if fill:
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
        ppr.insert_element_before(shd, *SUCC[1:])


def quote(t, size=10.5, space=3):
    """md 引用块：左侧深蓝竖线 + 浅蓝灰底，文字灰蓝，粗体藏蓝"""
    p = para(t, size, ea=EA, space=space, color=C_MINOR, indent=0.5)
    p.paragraph_format.right_indent = Cm(0.3)
    _pbdr(p, fill=FILL_ZEBRA)
    return p


def bullet(t, size=10.5):
    p = d.add_paragraph()
    r = p.add_run('• '); r.font.size = Pt(size); r.font.color.rgb = C_TITLE; set_fonts(r.element)
    add_runs(p, t, size)
    p.paragraph_format.left_indent = Cm(0.75); p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.28
    return p


def numbered(t, size=10.5):
    m = re.match(r'^(\d+)\. (.*)$', t)
    p = d.add_paragraph()
    r = p.add_run(m.group(1) + '. '); r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = C_TITLE; set_fonts(r.element)
    add_runs(p, m.group(2), size)
    p.paragraph_format.left_indent = Cm(0.75); p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.28
    return p


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def _tbl_borders(tb):
    tblpr = tb._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), BORDER); b.append(e)
    tblpr.insert_element_before(b, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')


def table(rows_):
    cells = [[c.strip() for c in re.split(r'(?<!\\)\|', r.strip().strip('|'))] for r in rows_ if not re.match(r'^\s*\|?\s*:?-{2,}', r)]
    if not cells: return
    ncol = max(len(r) for r in cells)
    tb = d.add_table(rows=len(cells), cols=ncol); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    _tbl_borders(tb)
    for i, row in enumerate(cells):
        for j in range(ncol):
            cell = tb.cell(i, j); cell.text = ''
            pr = cell.paragraphs[0]; pr.paragraph_format.space_after = Pt(1); pr.paragraph_format.line_spacing = 1.15
            txt = row[j] if j < len(row) else ''
            if i == 0:
                add_runs(pr, txt, 9, True, EA, C_WHITE, bold_color=None); _shade(cell, FILL_HEAD)
            else:
                add_runs(pr, txt, 9, False, EA, C_BODY)
                if i % 2 == 0: _shade(cell, FILL_ZEBRA)
    # 表头行跨页重复
    trpr = tb.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trpr.append(h)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def fig(path, cap, width=15.5):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width))
    para(cap, 9, ea=EA, space=10, color=C_MINOR, bold_color=C_EMPH)


CAPFIG = {'图 1': FIG1 + 'fig_truth_sigma_vs_z.png', '图 2': FIG1 + 'fig_env.png',
          '图 3': FIG1 + 'fig_llm.png', '图 4': FIG1 + 'fig_metric.png'}
E44CAP = ('图 5　内潮环境三轴扫描（84 km，同相位口径，末帧版）。(a) 500 km 隐藏锚：B 样条节点距 2/4/6 档的'
          '逐垂向模态远场功率保留比——第一模全档≈1.00，第二模 0.94→0.68→0.47；(b) 同一轴上 84 km 的三个可见读数'
          '（相对最轻档）：u/v 比值单调爬升报警，深层残余流缓降"变好看"，环带波功率（在本环境中只作展示、不参与判定）坠落；'
          '(c) 双谐波黏性八档与 (d) B 样条强度七档：残余流一路变好看，功率在 (d) 自 ×2 档后穿零反向，u/v 比值两轴'
          '均单调报警（黏性轴最低两档一处约 3% 回落，如实呈现）。')
CAP_ENV = '''图 1　探索环境四元组 (K, V, H, A)。K 是 Agent 可拧的旋钮，V 是公开读数（Agent 只看得到这一层），H 是保密判分（风驱侧为 MITgcm z 坐标独立求解器真值，内潮侧为 500 km 自参照锚，共用平底与静止态两条按构造零真值的通道），A 是九个审计算子。与文献中「把评测器藏起来」的现行做法相比，本环境的区别在于把 A 列为环境定义的第四项，并且用同一套算子同时审 V 与 H 自身——A3 只改盒宽就让隐藏判分侧的环带功率摆 7.24 倍，A4 修一行初值时钟让它从 4.58 变 7.89 MW。'''
CAP_AG = '''图 6　Agent 在与逻辑回归完全相同的信息条件下挑尺子（每局 40 个候选、七个无真值特征、不给真值，8 局）。(a) 逐局命中率：逻辑回归均值 0.675、低 B 单特征启发式 0.500、Agent 0.412、随机 0.283；Agent 8 局无一胜过逻辑回归。(b) 事后按 Agent 自述的抗刷率使用方向分组：说对方向的 3 局达 0.60（逼近线性模型），说反或没说清的 5 局回落到 0.30（随机水平）。结论是"稳定性而非能力"构成瓶颈；n=8、单模型、探索性，不作统计推断。'''
CAP_ASYM = '''图 5　两关的不对称。(a) 同一把 u/v 尺子在 0°/45°/90° 三个强迫方向下的两关得分：排序力 A 从 0.92 经 0.97 掉到 −0.17（变号），抗刷率 B 则在 1.00/0.25/0.89 之间摆动。(b) 四档地形陡度 rx0 两两比较：「B≥0.9 的候选集合」的 Jaccard 为 0.82–0.97，B 的候选间秩相关 0.73–0.95，预注册赌它掉到 0.5 以下（灰线）被推翻。结论：抗刷率是尺子自身的结构性质，对地形稳健；排序力是尺子与特定物理场景的匹配度，对强迫方向敏感。'''
CAP_TF = '''图 4　无真值预测器的命中率曲线（纬向风上训练，经向风上评估）。横轴是七个无真值特征经一次拟合的逻辑回归给出的预测分（对数轴），纵轴是该箱内真正通过排序力判据的比例，虚线为随机基率 0.2797216441207076。命中率从 0.037 升到峰值 0.8306，随后在最高分一箱回落到 0.6906；只取最极端的 0.5%（77 个）时跌至 0.052。跨环境 AUC = 0.8706。预测器在中段确有判别力，但在极端处与它所审判的那些代理指标一样失效——可操作的读法是取中段而非顶端。'''
CAP_PAR = '''图 3　排序力 A 与抗刷率 B 的双目标平面（风驱环境全部 15376 个候选）。(a) 灰点为全体候选，蓝点为两关同时通过者（204 个），橙线为 Pareto 前沿（5 点），紫星为乌托邦角 (1,1)——实测为空；虚线/点线为本文阈值 A≥0.7 / B≥0.9。knee 为前沿上到两端连线垂距最大的折中点，它在风驱环境两关全过，却不在跨内潮的 13 个交集里。(b) 阈值敏感性：A、B 阈值组合下的通过数，蓝框为本文口径。两关 Spearman = −0.674，两关同过 204 个而独立期望 2159 个（贫化 10.6 倍）。'''
CAP_KB = '''图 2　kill board：四把代表尺子 + 隐藏真值对照行 × 九项审计 + R 对照列。A9 为 Beckmann & Haidvogel (1993) 静止态海山检验；R 列为领域现行的 LLM 评审判据（内部逻辑/量纲/物理范围/文献接地）——四把尺子全过 R，而 A1–A9 逐项检验下没有一把能全部通过。† 残余流行：风驱侧为 deep_rms_800，内潮侧为 deep_dc_rms。行：环带波功率（我们论文头条）、深层残余流、深水温度（内潮穷举三轴满分冠军，实为 ≈4 °C 的深水温度本身）、u/v 比值（风驱穷举冠军）、隐藏真值（对照）。每格判定按 KILLBOARD.json 的 rules 字段判读，格内为关键数与出处；python scorecard.py 一键重生成。'''
CAP_E52 = '''图 2　转向经向风后的判据得分与交集计数。(a) u/v 与 v/u 在纬向风、经向风下的排序力（−ρ，对隐藏技巧评分，越高越好）与配对抗刷分率；虚线为判据阈值 A 0.7 / B 0.9。(b) 三个环境各自通过判据的候选数与两两、三重交集（对数轴，风驱空间 15376 / 内潮 19600）；竖线为独立随机判据的期望——三重交集 0 与期望 0.71 无差别，纬向∩内潮与经向∩内潮贫化、纬向∩经向富集 16 倍。'''
CAP_E44 = '''图 3　内潮三轴扫描（84 km，同相位口径，末帧版；(b) 的 m=5、(c) 的 5e7/2e8、(d) 的 ×½/×2/×8 为预注册后加密臂）。(a) 500 km 自参照模态锚：第一模全档≈1.00，第二模 0.94→0.68→0.47；(b–d) 三条旋钮轴上的可见读数（相对基准档）：u/v 单调报警、残余流变好看、环带功率坠落并在 (d) 穿零——但见图 1 后续审计。'''

V6 = 'v6' in SRC
IS_V5 = 'v5' in SRC


def body_line(ln):
    """普通行：列表 / 段落"""
    if re.match(r'^[-*] ', ln): bullet(ln[2:]); return
    if re.match(r'^\d+\. ', ln): numbered(ln); return
    para(ln, 10.5)


lines = md.split('\n')
i = 0
in_code = False; code_buf = []
while i < len(lines):
    ln = lines[i].rstrip()
    if ln.startswith('```'):
        if in_code:
            p = d.add_paragraph(); r = p.add_run('\n'.join(code_buf)); r.font.size = Pt(9)
            set_fonts(r.element, MONO, EA); r.font.color.rgb = C_EMPH
            p.paragraph_format.left_indent = Cm(0.5); _pbdr(p, fill=FILL_ZEBRA); p.paragraph_format.space_after = Pt(6)
            code_buf = []
        in_code = not in_code; i += 1; continue
    if in_code: code_buf.append(lines[i]); i += 1; continue
    if not ln: i += 1; continue
    if ln.startswith('> '):
        if ('v4' in SRC or IS_V5 or V6) and not ln.startswith('> 初赛定稿'):
            quote(ln[2:])
        i += 1; continue
    if ln.startswith('|'):
        blk = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            blk.append(lines[i]); i += 1
        table(blk); continue
    if ln == '---': i += 1; continue
    m = re.match(r'^(#{1,3}) (.+)$', ln)
    if m:
        lvl = len(m.group(1)); t = m.group(2)
        # v6：图 4/5/6 放在 4.9 之前（4.6–4.8 正文引用处之后）
        if V6 and lvl == 2 and t.startswith('4.9 '):
            fig(FIG2 + 'fig_truthfree.png', CAP_TF, 14.0); fig(FIG2 + 'fig_asym.png', CAP_ASYM, 15.0); fig(FIG2 + 'fig_agent_picks.png', CAP_AG, 15.0)
        if lvl == 1 and t.startswith('谁来给尺子打分'):
            p = para(t, 20, True, EA_HEAD, 4, WD_ALIGN_PARAGRAPH.CENTER, color=C_TITLE, bold_color=None)
            p.paragraph_format.space_before = Pt(6)
        elif lvl == 2 and t.startswith('复赛'):
            p = para(t, 12, True, EA_HEAD, 4, WD_ALIGN_PARAGRAPH.CENTER, color=C_MINOR, bold_color=None)
            p.paragraph_format.space_after = Pt(10)
            # 标题下分隔线
            ppr = p._p.get_or_add_pPr(); bdr = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
            bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '12'); bt.set(qn('w:space'), '6'); bt.set(qn('w:color'), FILL_HEAD)
            bdr.append(bt); ppr.insert_element_before(bdr, 'w:shd', 'w:tabs', 'w:spacing', 'w:ind', 'w:jc', 'w:rPr', 'w:sectPr', 'w:pPrChange')
        else:
            heading(t, lvl)
        if lvl == 1 and t.startswith('第四问') and not IS_V5 and not V6:
            if 'v4' in SRC:
                fig(FIG2 + 'fig_killboard_print.png', CAP_KB, 14.5)
                fig(FIG2 + 'fig_e52.png', CAP_E52, 15.0)
                fig(FIG2 + 'fig_e44_main.png', CAP_E44, 15.5)
            else:
                fig(FIG2 + 'fig_e44_main.png', E44CAP, 15.5)
        if V6 and lvl == 2 and t.startswith('2.1 '):
            fig(FIG2 + 'fig_env.png', CAP_ENV, 15.5)
        i += 1; continue
    cm = re.match(r'^(图 \d)　', ln)
    if cm and cm.group(1) in CAPFIG:
        pth = CAPFIG.pop(cm.group(1))
        pp = d.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.keep_with_next = True
        pp.add_run().add_picture(pth, width=Cm(12.5))
        para(ln, 9, ea=EA, space=8, color=C_MINOR)
        i += 1; continue
    body_line(ln)
    if (IS_V5 or V6) and ln.startswith('**参照系**'):
        fig(FIG2 + 'fig_killboard_print.png', CAP_KB, 14.5)
        if V6: fig(FIG2 + 'fig_pareto.png', CAP_PAR, 15.5)
        if not V6:
            fig(FIG2 + 'fig_e52.png', CAP_E52, 15.0); fig(FIG2 + 'fig_e44_main.png', CAP_E44, 15.5)
    i += 1

# ---- 附录合订 (v5/v6) ----
if IS_V5 or V6:
    def page_break(): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def md_block(text, plain_headings=False):
        ls = text.split('\n'); k = 0
        while k < len(ls):
            ln = ls[k].rstrip()
            if not ln or ln == '---': k += 1; continue
            mm = re.match(r'^(#{1,3}) (.+)$', ln)
            if mm: heading(mm.group(2), len(mm.group(1))); k += 1; continue
            if plain_headings:  # 初赛定稿提取文本没有 # 标记，按「一、」「1.1 」编号识别标题
                if re.match(r'^[一二三四五六七八九十]+、\S{1,20}$', ln): heading(ln, 2); k += 1; continue
                if re.match(r'^\d\.\d+ \S.{0,30}$', ln): heading(ln, 3); k += 1; continue
            if ln.startswith('|'):
                blk = []
                while k < len(ls) and ls[k].strip().startswith('|'): blk.append(ls[k]); k += 1
                table(blk); continue
            if ln.startswith('> '): quote(ln[2:], 10); k += 1; continue
            body_line(ln); k += 1
    page_break(); heading('附录 A · 初赛定稿正文（原文，2026-08-16 提交版）', 1, 0)
    md_block(open(F + '初赛终稿_提取.txt').read(), plain_headings=True)
    page_break(); md_block(open(F + '附录B_实验表.md').read())
    page_break(); md_block(open(F + '附录C_ARENA规格.md').read())
    # 页码（页脚居中，灰蓝小字）
    for sec in d.sections:
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        for tag in ('begin', 'instr', 'end'):
            if tag == 'instr':
                el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = 'PAGE'
            else:
                el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), tag)
            run._r.append(el)
        run.font.size = Pt(9); run.font.color.rgb = C_MINOR; set_fonts(run.element)
d.save(F + OUT)
print('docx saved')
