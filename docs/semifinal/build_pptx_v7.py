# -*- coding: utf-8 -*-
"""按规格 JSON 生成 16:9 可编辑 PPTX（微软雅黑、大字号、窄边距）。
改自算法赛题队 scripts/build_pptx.py（v7 版）：
  - 封面长标题自动降到 36pt 并加高标题框
  - kind=split：左要点 / 右图；kind=image：上图 / 下两行说明；kind=bullets：要点（可带 bignum / quote）
  - 要点字号在 24→20pt 之间按版心高度选最大能放下的，绝不低于 20pt；放不下就打印 WARNING 让人删字
    python build_pptx_v7.py spec.json 输出.pptx --figs 图目录
"""
import json, os, sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt
from lxml import etree

FONT = "微软雅黑"
# 官方模板配色（AI for Research 初赛模板 docx 提取）
BLUE = RGBColor(0x1F, 0x4D, 0x78)      # 标题色 / 标题条 / 封面底
NAVY = RGBColor(0x0B, 0x25, 0x45)      # 强调色 / 封面下条 / 正文加粗强调
GRAY = RGBColor(0x5A, 0x65, 0x73)      # 次要文字（页码、说明）
BLUE_L = RGBColor(0xF4, 0xF6, 0xF9)    # 表格底纹 / 引用框、数字框浅底
ORANGE = NAVY                          # 正文强调沿用变量名，颜色改为官方强调色
DARK = RGBColor(0x26, 0x2A, 0x2E)      # 正文
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TINT = RGBColor(0xC9, 0xD6, 0xE3)      # 深底上的次要文字（浅蓝灰）
TINT_L = RGBColor(0xF4, 0xF6, 0xF9)    # 深底上的正文（近白）
BG = BLUE_L
W, H = Inches(13.333), Inches(7.5)
M = Inches(0.42)
BAR = Inches(1.02)
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
WARN = []


def set_font(run, size, bold=False, color=DARK, name=FONT):
    f = run.font
    f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        for old in rPr.findall(NS + tag):
            rPr.remove(old)
        el = etree.SubElement(rPr, NS + tag)
        el.set("typeface", name)


def box(slide, x, y, w, h, fill=None, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    for i, (t, sz, bold, col) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        p.line_spacing = 1.22
        r = p.add_run(); r.text = t
        set_font(r, sz, bold, col)
    return tb


def emph_runs(paragraph, s, keys, size, color=DARK):
    import re
    if not keys:
        r = paragraph.add_run(); r.text = s; set_font(r, size, False, color); return
    pat = "|".join(re.escape(k) for k in sorted(keys, key=len, reverse=True) if k)
    parts = re.split(f"({pat})", s)
    for p in parts:
        if not p: continue
        r = paragraph.add_run(); r.text = p
        set_font(r, size, p in keys, ORANGE if p in keys else color)


def units(s):
    """中文按 1 字宽、拉丁按 0.55 字宽"""
    return sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in s)


def wrapped_lines(lines, size, width_in):
    cap = max(8, int(width_in * 72 / size) - 1)
    n = 0
    for ln in lines:
        cost = units(ln) + 1.5   # 含项目符号
        n += max(1, -(-int(cost * 100) // (cap * 100)))
    return max(1, n)


def bullets_height(lines, size, width_in):
    n = wrapped_lines(lines, size, width_in)
    return n * (size * 1.3 + 11) / 72 + 0.20


def add_bullets(slide, x, y, w, h, lines, keys, sizes=(24, 23, 22, 21, 20), anchor=MSO_ANCHOR.TOP, tag=""):
    w_in = w / 914400 - 0.1
    h_in = h / 914400
    sz = sizes[-1]
    for s_ in sizes:
        if bullets_height(lines, s_, w_in) <= h_in:
            sz = s_; break
    need = bullets_height(lines, sz, w_in)
    if need > h_in + 0.02:
        WARN.append(f"{tag}: 要点在 {sz}pt 下估算需 {need:.2f}in，只有 {h_in:.2f}in，请删字")
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(11); p.line_spacing = 1.3
        r = p.add_run(); r.text = "▪ "; set_font(r, sz, True, BLUE)
        emph_runs(p, ln.strip().lstrip("-•▪ "), keys, sz)
    return sz, need


def pic_fit(slide, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih)
    nw, nh = int(iw * sc), int(ih * sc)
    return slide.shapes.add_picture(path, int(x + (w - nw) / 2), int(y + (h - nh) / 2), nw, nh)


def keys_of(b):
    return [k for k in (b.get("emph") or "").replace("，", "、").split("、") if k.strip()]


def lines_of(b):
    return [x for x in b["content"].strip().splitlines() if x.strip()]


def build(spec, out, figs):
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    N = len(spec["deck"])
    for pg in spec["deck"]:
        s = prs.slides.add_slide(blank)
        kind = pg.get("kind", "bullets")
        blocks = pg.get("blocks", [])
        tag = f"第 {pg['n']} 页"
        if kind == "cover":
            box(s, 0, 0, W, H, BLUE)
            box(s, 0, int(H * 0.78), W, int(H * 0.22), NAVY)
            tl = pg["title"].split("\n")   # 允许用换行控制封面标题在哪儿折行，字不变
            tsz = 42 if max(units(x) for x in tl) <= 20 else 36
            text(s, M + Inches(0.5), Inches(1.05), W - 2 * M - Inches(1.0), Inches(1.55),
                 [(x, tsz, True, WHITE) for x in tl], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, space=2)
            if pg.get("subtitle"):
                text(s, M + Inches(1.0), Inches(2.68), W - 2 * M - Inches(2.0), Inches(0.5),
                     [(pg["subtitle"], 17, False, TINT)], PP_ALIGN.CENTER)
            q = next((b for b in blocks if b["type"] == "quote"), None)
            if q:
                text(s, M + Inches(0.6), Inches(3.30), W - 2 * M - Inches(1.2), Inches(1.05),
                     [(q["content"].strip(), 21, False, TINT_L)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            big = next((b for b in blocks if b["type"] == "bignum"), None)
            if big:
                num, _, desc = big["content"].partition("||")
                text(s, M, Inches(4.40), W - 2 * M, Inches(0.85),
                     [(num.strip(), 48, True, WHITE)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
                text(s, M + Inches(1.2), Inches(5.22), W - 2 * M - Inches(2.4), Inches(0.55),
                     [(desc.strip(), 16, False, TINT)], PP_ALIGN.CENTER)
            bl = next((b for b in blocks if b["type"] == "bullets"), None)
            if bl:
                text(s, M, Inches(6.15), W - 2 * M, Inches(0.9),
                     [(x.strip(), 16, False, TINT_L)
                      for x in lines_of(bl)], PP_ALIGN.CENTER, space=3)
            if pg.get("notes"):
                s.notes_slide.notes_text_frame.text = pg["notes"]
            continue

        # 标题条：标题长就占满整条（不放右侧小字）
        box(s, 0, 0, W, BAR, BLUE)
        box(s, 0, BAR, W, Inches(0.06), NAVY)   # 标题条下沿一道藏蓝细线
        tu = units(pg["title"])
        tsz = 31 if tu <= 24 else 30
        full = (not pg.get("subtitle")) or tu > 21
        tw = W - 2 * M if full else W - 2 * M - Inches(3.0)
        if tu > int(tw / 914400 * 72 / tsz) - 1:
            WARN.append(f"{tag}: 标题 {tu:.0f} 字宽，{tsz}pt 下可能折成两行")
        text(s, M, Inches(0.14), tw, Inches(0.76), [(pg["title"], tsz, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
        if pg.get("subtitle") and not full:
            text(s, W - M - Inches(3.0), Inches(0.22), Inches(3.0), Inches(0.6),
                 [(pg["subtitle"], 14, False, TINT)], PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)

        top = BAR + Inches(0.22); avail = H - top - Inches(0.30)
        body_w = W - 2 * M
        imgs = [b for b in blocks if b["type"] == "image"]
        buls = [b for b in blocks if b["type"] == "bullets"]
        quotes = [b for b in blocks if b["type"] == "quote"]
        bigs = [b for b in blocks if b["type"] == "bignum"]
        img_path = os.path.join(figs, imgs[0]["content"].strip()) if imgs else None
        if img_path and not os.path.exists(img_path):
            WARN.append(f"{tag}: 找不到图 {img_path}"); img_path = None

        if kind == "split" and img_path and buls:
            left_w = Inches(float(pg.get("left_w", 5.7)))
            gap = Inches(0.25)
            b = buls[0]
            add_bullets(s, M, top, left_w, avail, lines_of(b), keys_of(b), anchor=MSO_ANCHOR.MIDDLE, tag=tag)
            pic_fit(s, img_path, M + left_w + gap, top, body_w - left_w - gap, avail)
        elif kind == "image" and img_path:
            y_img_h = avail
            if buls:
                b = buls[0]; ls = lines_of(b)
                sz = 22 if len(ls) <= 2 else 21
                bh = Inches(bullets_height(ls, sz, body_w / 914400 - 0.1))
                y_img_h = avail - bh
                add_bullets(s, M, top + y_img_h, body_w, bh, ls, keys_of(b), sizes=(sz, 21, 20), tag=tag)
            pic_fit(s, img_path, M, top, body_w, y_img_h)
        else:  # bullets（可带 bignum 在上、quote 在下）
            y = top; rest = avail
            if bigs:
                num, _, desc = bigs[0]["content"].partition("||")
                bh = Inches(1.45)
                box(s, M, y, body_w, bh, BLUE_L)
                box(s, M, y, Inches(0.09), bh, BLUE)
                text(s, M, y + Inches(0.05), body_w, int(bh * 0.60),
                     [(num.strip(), 48, True, BLUE)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
                text(s, M, y + int(bh * 0.60), body_w, int(bh * 0.40),
                     [(desc.strip(), 20, False, DARK)], PP_ALIGN.CENTER, MSO_ANCHOR.TOP)
                y += bh + Inches(0.15); rest -= bh + Inches(0.15)
            qh = 0
            if quotes:
                q = quotes[0]
                nl = wrapped_lines([q["content"].strip()], 20, body_w / 914400 - 0.6)
                qh = Inches(nl * 20 * 1.3 / 72 + 0.45)
                rest -= qh + Inches(0.15)
            if buls:
                b = buls[0]
                add_bullets(s, M, y, body_w, rest, lines_of(b), keys_of(b),
                            anchor=MSO_ANCHOR.MIDDLE, tag=tag)
                y += rest + Inches(0.15)
            if quotes:
                q = quotes[0]
                box(s, M, y, body_w, qh, BLUE_L)
                box(s, M, y, Inches(0.09), qh, BLUE)   # 引用框左侧深蓝竖条
                tb = s.shapes.add_textbox(M + Inches(0.3), y, body_w - Inches(0.55), qh)
                tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = tf.paragraphs[0]; p.line_spacing = 1.25
                emph_runs(p, q["content"].strip(), keys_of(q), 20, BLUE)
        text(s, W - M - Inches(1.0), H - Inches(0.36), Inches(1.0), Inches(0.3),
             [(f"{pg['n']} / {N}", 11, False, GRAY)], PP_ALIGN.RIGHT)
        if pg.get("notes"):
            s.notes_slide.notes_text_frame.text = pg["notes"]
    prs.save(out)
    print(f"写出 {out}  {N} 页")
    for w in WARN:
        print("WARNING", w)


if __name__ == "__main__":
    figs = "demo"
    a = sys.argv[1:]
    if "--figs" in a:
        k = a.index("--figs"); figs = a[k + 1]; a = a[:k] + a[k + 2:]
    build(json.load(open(a[0], encoding="utf-8")), a[1], figs)
