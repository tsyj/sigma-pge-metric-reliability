#!/home/xinyuan/anaconda3/envs/numpy1/bin/python
# -*- coding: utf-8 -*-
"""复赛 PPT v5（6 页, 3 分钟）"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os, json
BLUE=RGBColor(0x3C,0x54,0x88); DARK=RGBColor(0x1a,0x1a,0x1a); GRAY=RGBColor(0x66,0x66,0x66); WHITE=RGBColor(0xff,0xff,0xff)
FIG1='/home/xinyuan/比赛/赛道三赛题二/figs/'; FIG2='/data/xinyuan/GOAI_ai4s_env/e44/figs/'; FONT='Noto Sans CJK SC'
E54=json.load(open('/data/xinyuan/GOAI_ai4s_env/e44/agent/E54_NAMESWAP.json'))['conditions']
TAG=os.environ.get('FUSAI_TAG','fusai-v5 @ 〔SHA〕')
E55=os.environ.get('E55_LINE','45° 风：结果回填中')
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
def txt(s,x,y,w,h,size,color=DARK,bold=False,align=PP_ALIGN.LEFT,line=None):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame; tb.word_wrap=True
    for i,t in enumerate(line if isinstance(line,list) else [line or '']):
        p=tb.paragraphs[0] if i==0 else tb.add_paragraph(); p.text=t; p.alignment=align
        for r in p.runs: r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold
def head(s,title,sub=None):
    txt(s,0.55,0.25,12.2,0.9,28,BLUE,True,line=[title])
    if sub: txt(s,0.55,1.0,12.2,0.6,15,GRAY,line=[sub])
def pic(s,path,x,y,w):
    if os.path.exists(path): s.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w))
# p1 封面：一句话
s=prs.slides.add_slide(blank); r=s.shapes.add_shape(1,0,0,prs.slide_width,prs.slide_height); r.fill.solid(); r.fill.fore_color.rgb=BLUE; r.line.fill.background()
txt(s,1.0,1.6,11.3,1.3,42,WHITE,True,line=['谁来给尺子打分'])
txt(s,1.0,2.8,11.3,0.7,18,WHITE,line=['σ 坐标伪流环境中的代理指标可靠性探索 · 复赛 · 队伍：虎虎'])
txt(s,1.0,4.4,11.3,2.0,24,WHITE,True,line=['三万五千把尺子、三种风向、八项审计，一把不剩——','所以我们交的不是尺子，是给尺子打分的协议。'])
# p2 环境 = (K,V,H,A) + 问题小图
s=prs.slides.add_slide(blank); head(s,'把真值藏起来，做一个给评价指标打分的环境','陡海山上模式算出假流：真值 9.2 cm/s，σ 坐标 100 cm/s。评价"治好没有"的读数，平时没有真值可对')
pic(s,FIG1+'fig_truth_sigma_vs_z.png',0.5,1.75,5.6)
txt(s,6.4,1.9,6.6,5,19,DARK,line=['环境 = (K, V, H, A)','K 旋钮可拧：黏性、B 样条、地形、积分','V 读数可见：流速、残余流、比值、功率','H 判分隐藏：MITgcm 独立解 / 500 km 模态锚 / 平底零通道','A 八把审计的刀：换平底、零通道、盒宽、相位、盖章预测、转风向、匿名换名、真值锚','','三种驱动共用一套接口；每条算例 md5 留痕'])
# p3 kill board (slide 版)
s=prs.slides.add_slide(blank); head(s,'主图：四把尺子 × 八项审计——没有一行全绿','对照行 = 隐藏真值自己走一遍：该放过的放过，该杀的杀。每格一个数、一个出处，python scorecard.py 一键重生成')
pic(s,FIG2+'fig_killboard_slide.png',0.55,1.65,12.2)
# p4 主结果 E52 图通栏
s=prs.slides.add_slide(blank); head(s,'主结果：风向一转（45° / 90°），排序力与抗刷分从未同时落在同一把尺子上','90°：u/v 排序 −0.92→+0.17；转成 v/u 排序回来但抗刷 0.11。45°（预注册）：u/v 排序 −0.97 却抗刷 0.25')
pic(s,FIG2+'fig_e52.png',1.1,1.55,7.6)
txt(s,8.9,1.7,4.2,5.4,15.5,DARK,line=['• 三重交集 0，与随机期望 0.71 无差别——"0"不是证据；证据是贫化与死因','• 排序看强迫方向，抗刷看区域几何（x 周期、y 陆墙），只在纬向风下碰巧重合','• '+E55,'• 自家论文头条也死：环带功率只改盒宽摆 7.24×，修一行时钟 4.58→7.89——论文将更正'])
# p5 Agent 三角色（表）
s=prs.slides.add_slide(blank); head(s,'Agent 的三个角色：对手、评审、提议者——目前不是发现者','单模型、小 n，全部明写')
txt(s,0.9,2.0,11.6,4.8,21,DARK,line=[
 '对手  ｜ 目标"残余流最低" → 交出波功率只剩 30% 的方案；中性措辞终选相同  ｜ n=2 局 + 3 席随机',
 '评审  ｜ 真名：信坏尺子 %d/5、疑好尺子 %d/5 → 匿名 %d/5、%d/5  ｜ 名字本身是污染源'%(E54['C0_named']['Pnet_trusted'],E54['C0_named']['uv_gameable'],E54['C1_anon']['Pnet_trusted'],E54['C1_anon']['uv_gameable']),
 '提议者｜ 20 轮有效提议 0 通过，随机期望 0.77 概率也是 0  ｜ 与随机无差别',
 '','下一步：出题者-刷分者对抗环境（ARENA，规格已写，缓存库 430 条，每轮 ≤2 分钟）',
 '——先证明标量尺子不存在，再让 Agent 去找，才不是让它追一个已知不存在的目标'])
# p6 沉淀
s=prs.slides.add_slide(blank); head(s,'交的是协议，不是尺子','137 条算例一次完赛、md5 审计；3 条盖章预测被反驳、1 次指标降格、1 次论文更正、1 次环境 bug 声明')
txt(s,0.9,2.0,11.6,4.8,21,DARK,line=['• 三驱动一接口、三形态真值、A1–A8 审计脚本、kill board 一键重生成、预注册协议（含 45° 风盖章预测）',
 '• 换尺子直接跑；换模式实现三个函数（submit_run / cache_lookup / analyze）',
 '• 仓库 github.com/tsyj/sigma-pge-metric-reliability  '+TAG,
 '','三万五千把尺子、三种风向、八项审计，一把不剩——所以我们交的不是尺子，是给尺子打分的协议。'])
prs.save('/home/xinyuan/比赛/赛道三赛题二/复赛/方案说明PPT_v5.pptx'); print('saved PPT v5')
