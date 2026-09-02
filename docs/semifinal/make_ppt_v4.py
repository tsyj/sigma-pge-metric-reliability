#!/home/xinyuan/anaconda3/envs/numpy1/bin/python
# -*- coding: utf-8 -*-
"""复赛 PPT v4（方案 B, 7 页, 3 分钟）"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os, json
BLUE=RGBColor(0x3C,0x54,0x88); DARK=RGBColor(0x1a,0x1a,0x1a); GRAY=RGBColor(0x66,0x66,0x66); WHITE=RGBColor(0xff,0xff,0xff); AMBER=RGBColor(0xC4,0x6A,0x1F)
FIG1='/home/xinyuan/比赛/赛道三赛题二/figs/'; FIG2='/data/xinyuan/GOAI_ai4s_env/e44/figs/'; FONT='Noto Sans CJK SC'
E54P='/data/xinyuan/GOAI_ai4s_env/e44/agent/E54_NAMESWAP.json'
e54=json.load(open(E54P)) if os.path.exists(E54P) else None
def k(c,key):
    try: return '%d/%d'%(e54['conditions'][c][key],e54['conditions'][c]['n'])
    except Exception: return 'k/5'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
def txt(s,x,y,w,h,size,color=DARK,bold=False,align=PP_ALIGN.LEFT,line=None):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame; tb.word_wrap=True
    for i,t in enumerate(line if isinstance(line,list) else [line or '']):
        p=tb.paragraphs[0] if i==0 else tb.add_paragraph(); p.text=t; p.alignment=align
        for r in p.runs: r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold
def head(s,title,sub=None):
    txt(s,0.55,0.25,12.2,0.9,28,BLUE,True,line=[title])
    if sub: txt(s,0.55,1.0,12.2,0.6,15,GRAY,line=[sub])
def pic(s,path,x,y,w):
    if os.path.exists(path): s.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w))
# 1 封面 + 一句话
s=prs.slides.add_slide(blank); r=s.shapes.add_shape(1,0,0,prs.slide_width,prs.slide_height); r.fill.solid(); r.fill.fore_color.rgb=BLUE; r.line.fill.background()
txt(s,1.0,1.7,11.3,1.3,40,WHITE,True,line=['谁来给尺子打分'])
txt(s,1.0,2.9,11.3,0.7,19,WHITE,line=['σ 坐标伪流环境中的代理指标可靠性探索 · 复赛 · 队伍：虎虎'])
txt(s,1.0,4.3,11.3,2.4,17,WHITE,line=['尺子的可靠性不是标量的性质，是审计协议的性质。',
 '我们把真值藏起来，让每把尺子都能被打分——作者自己的论文头条、有名字的大模型评审、两万选一的穷举冠军，',
 '都在"看起来合理"上栽了；唯一活到最后的 u/v 比值，被审计一路收窄，换到第三种强迫也死了。'])
# 2 问题
s=prs.slides.add_slide(blank); head(s,'问题：伪流治理的每个读数，都是一把没打过分的尺子','同一座海山：z 坐标真值 9.2 cm/s，σ 坐标 100 cm/s。而评价"治理好不好"的读数，平时没有真值可对')
pic(s,FIG1+'fig_truth_sigma_vs_z.png',2.6,1.75,8.1)
# 3 环境 = (K,V,H,A)
s=prs.slides.add_slide(blank); head(s,'环境 = (K 旋钮, V 可见读数, H 隐藏判分, A 审计算子)','三种驱动共用一套接口；真值三种形态（MITgcm 独立解 / 500 km 模态锚 / 平底零通道）；审计算子是一等公民')
txt(s,0.9,2.0,11.6,5,17,DARK,line=[
 'A1 配对抗刷：同参数换到平底，读数会不会变好看 —— 杀了残余流类（抗刷率 0.000）',
 'A2 零真值通道：真值恒为零处读数是否为零 —— u/v 在平底发散',
 'A3 盒宽扰动：只改盒宽读数动多少 —— 环带功率摆 7.24×',
 'A4 时钟/相位扰动：改一行初值时钟、±T/8 —— 环带功率 4.58→7.89；末帧类读数 ±30–66%',
 'A5 盖章预测：先写后跑 —— "u/v 会失效"被反驳，"守门会退化"落空',
 'A6 跨环境（第三强迫）：换驱动方向还过不过 —— 两环境交集 13 把全部失效',
 'A7 匿名/换名：摘掉名字评审判断变不变 —— 有名字时坏尺子被信、好尺子被疑',
 'A8 真值锚：与隐藏真值的秩相关 —— 暴露内潮真值单侧，残余流比 u/v 更贴锚'])
# 4 kill board
s=prs.slides.add_slide(blank); head(s,'主结果：四把尺子 × 八个审计——没有一行全绿','每格一个判定、一个数、一个出处文件；python scorecard.py 一键重生成')
pic(s,FIG2+'fig_killboard.png',0.5,1.65,12.3)
# 5 头条之死 + E52
s=prs.slides.add_slide(blank); head(s,'两次亲手证伪：自家论文头条，和自家穷举冠军','左：环带"相干辐射功率"量的是域共振。右：换到第三种强迫，两环境交集 13 把尺子全部失效，三重交集 0')
txt(s,0.5,1.75,6.0,5.0,14.5,DARK,line=['环带波功率（上过我们论文头条 "只保留 11.1%"）：','• 只改盒宽：摆 7.24×，正压响应只差 1.5%','• 修一行初值时钟：4.58 → 7.89 MW','• 换 500 km 干净口径："物理最差 11%" 0.52 → 0.94','→ 降为环境内游戏分数；论文将作更正','','u/v 比值（15376 选 1 的风驱冠军）：','• 内潮三轴单调报警，预注册"会失效"被反驳','• 但：真值单侧下残余流更贴锚（+0.90 vs −0.70）','• 相位 ±30–66%（窗均 4%）、平底发散、盒宽绝对值变质','• 经向风：排序力 −0.92 → +0.17，三重交集 0'])
pic(s,FIG2+'fig_e52.png',6.6,2.0,6.5)
# 6 LLM 三角色
s=prs.slides.add_slide(blank); head(s,'Agent 的三个角色：对手、评审、提议者——目前不是发现者','单模型、小 n，全部明写')
txt(s,0.9,2.0,11.6,5,16.5,DARK,line=[
 '对手（陷阱局，n=2 措辞 + 3 席随机）：目标"把残余流降到最低"→ 交出波功率 −70% 的方案；可见 u/v 一路报警没用；',
 '     换成中性措辞终选一模一样，还把 u/v 标"误导"、把环带功率标"可信"。随机基线 2/3 席掉得更深。',
 '评审（换名局 ×5）：真名 → 信环带功率 '+k('C0_named','Pnet_trusted')+'、疑 u/v '+k('C0_named','uv_gameable')+'；匿名 → 信环带功率 '+k('C1_anon','Pnet_trusted')+'、疑 u/v '+k('C1_anon','uv_gameable')+'；',
 '     互换名 → 信"环带功率"(实为 u/v) '+k('C2_swapped','Pnet_trusted')+'。名字本身是污染源。',
 '提议者（初赛 E43）：30 轮 0 通过；按穷举基率 1.33%，30 抽 0 中的概率 0.67——与随机无差别。',
 '',
 '结论：Agent 现在是对手和被试；缺口具体到"构造零真值子空间"。下一步：出题者-刷分者对抗环境（ARENA，规格已写）。'])
# 7 负结果说明什么 + 仓库
s=prs.slides.add_slide(blank); head(s,'负结果说明什么：可靠性是协议的性质','15376–19600 候选 × 三种强迫 × 八个算子：没有普适的标量尺子——但每一格都能追到一条审计、一个数、一个文件')
txt(s,0.9,2.1,11.6,4.8,17.5,DARK,line=[
 '• 一把比值尺子要同时满足两件独立的事：排序力（真信号钉住一端）与抗刷分（换平底不变好看）',
 '• 前者由真信号住在哪个分量决定（强迫方向），后者由区域几何决定（x 周期、y 陆墙）',
 '• 纬向强迫下两者巧合落在同一对 u/v；经向下分家——没有一个比值同时满足。最粗代理排序 ρ≈−0.99 却完全不抗刷',
 '',
 '• 沉淀为可接续的环境包：三驱动一接口、三形态真值、A1–A8 审计脚本、kill board 一键重生成、预注册协议',
 '• 88 条算例一次完赛、md5 审计、健康门全绿；1 条预测被反驳、1 次指标降格、1 次论文更正、1 次环境 bug 声明',
 '• github.com/tsyj/sigma-pge-metric-reliability（复赛 tag 冻结）'])
prs.save('/home/xinyuan/比赛/赛道三赛题二/复赛/方案说明PPT_v4.pptx'); print('saved PPT v4')
